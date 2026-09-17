"""Build a local research GIS layer from user-authorized GESIS downloads.

Requires pandas and openpyxl for read-only extraction. Outputs, including all
derived data, stay under git-ignored tmp/. Never bundle them with the package.
"""
import csv
import hashlib
import json
import pathlib
import re
import xml.etree.ElementTree as ET

import pandas as pd
from openpyxl import load_workbook

ROOT = pathlib.Path(__file__).resolve().parents[1]
OUT = ROOT / 'tmp/industrial-research'
NS = {'k': 'http://www.opengis.net/kml/2.2'}
DOI = 'https://doi.org/10.7802/2762'


def coordinates(element):
    return [[float(n) for n in token.split(',')[:2]]
            for token in element.text.split()]


def main():
    workbook = OUT / 'Vietnam_SEZ_Metadata.xlsx'
    assert hashlib.md5(workbook.read_bytes()).hexdigest() == '0ecc15a8262f518c37d5c1a34067cf2d'
    df = pd.read_excel(workbook, sheet_name='SEZ metadata')
    df.columns = [re.sub(r'\s+', ' ', x).strip() for x in df.columns]
    df['sez_id'] = df['SEZ code'].map(lambda x: str(int(x)).zfill(7))
    assert len(df) == 611 and df.sez_id.is_unique
    # Preserve actual hyperlink targets, not merely Excel's display word Source.
    ws = load_workbook(workbook, read_only=False, data_only=True)['SEZ metadata']
    df['metadata_source_hyperlink'] = [ws.cell(i, 16).hyperlink.target
        if ws.cell(i, 16).hyperlink else None for i in range(2, len(df) + 2)]
    tree = ET.parse(OUT / 'Vietnam_SEZ_Polygons.kml')
    points, polygons, geometry_issues = {}, [], []
    for p in tree.findall('.//k:Placemark', NS):
        name = p.findtext('k:name', namespaces=NS) or ''
        point = p.find('k:Point/k:coordinates', NS)
        if point is not None:
            assert name not in points, 'Duplicate point ID: ' + name
            points[name] = coordinates(point)[0]
        parts = []
        for poly in p.findall('.//k:Polygon', NS):
            rings = [coordinates(r) for r in poly.findall('.//k:LinearRing/k:coordinates', NS)]
            if not rings or not all(len(r) >= 4 and r[0] == r[-1] for r in rings):
                geometry_issues.append({'label': name, 'ring_lengths': [len(r) for r in rings],
                                        'issue': 'empty_or_invalid_source_ring_not_exported'})
                continue
            parts.append(rings)
        if parts:
            match = re.match(r'(\d{7})\s*[;,]', name)
            sid = match[1] if match else name[:7]
            polygons.append({'type': 'Feature', 'properties': {
                'sez_id': sid, 'observation_label': name, 'source_url': DOI,
                'geometry_meaning': 'satellite_observed_built_up_area_not_legal_boundary'},
                'geometry': {'type': 'MultiPolygon', 'coordinates': parts}})
    assert set(points) == set(df.sez_id), 'Metadata and KML IDs differ'
    df['longitude'] = df.sez_id.map(lambda x: points[x][0])
    df['latitude'] = df.sez_id.map(lambda x: points[x][1])
    assert df.longitude.between(102, 110).all() and df.latitude.between(8, 24).all()
    df['source_url'] = DOI
    df['source_version'] = '1.0.0'
    df['coverage_period'] = '1991–2022'
    df['location_accuracy'] = 'research_source_point'
    df['attribute_review'] = df['SEZ name'].isna().map({True: 'missing_source_name_and_type', False: 'source_attributes_retained'})
    df.to_csv(OUT / 'sez-research-sites.csv', index=False)
    records = json.loads(df.to_json(orient='records', force_ascii=False))
    features = [{'type': 'Feature', 'properties': r,
                 'geometry': {'type': 'Point', 'coordinates': points[r['sez_id']]}} for r in records]
    for filename, rows in [('sez-research-sites.geojson', features), ('sez-built-up-history.geojson', polygons)]:
        (OUT / filename).write_text(json.dumps({'type': 'FeatureCollection', 'features': rows}, ensure_ascii=False))
    counts = df['Type of Zone'].fillna('Missing in source').value_counts().to_dict()
    audit = {'metadata_rows': len(df), 'unique_point_ids': len(points),
             'metadata_without_point': 0, 'points_without_metadata': 0,
             'missing_names': int(df['SEZ name'].isna().sum()),
             'source_polygon_issues': geometry_issues,
             'type_counts': counts, 'built_up_observations': len(polygons),
             'polygon_ids_without_metadata': sorted({p['properties']['sez_id'] for p in polygons} - set(points)),
             'files': {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                       for p in sorted(OUT.iterdir()) if p.suffix in ('.xlsx', '.kml', '.dta')}}
    (OUT / 'research-audit.json').write_text(json.dumps(audit, indent=2, ensure_ascii=False))
    (OUT / 'README.md').write_text('''# Vietnam SEZ research data

User-authorized scientific research use only. Raw files and derived layers must
remain local and must not be committed, published, or redistributed without the
permission required by https://data.gesis.org/sharing/#!TermsOfUse.
Delete the research data when use is complete. Cite and report resulting
publications as required by those terms.

Citation: Tafese, T., Lay, J., & Tran, V. (2025). Vietnam Special Economic Zone
Satellite Imagery Dataset (Version 1.0.0). GIGA. https://doi.org/10.7802/2762

All 611 metadata IDs have KML points. The historical source covers 1991–2022,
not a complete September 2026 register. One original metadata row lacks a name
and zone type; it is retained with its ID and source point. KML polygon snapshots
represent observed built-up areas, not legal industrial-park boundaries.
SEZ IDs are seven-character strings; leading zeros must be preserved.
''')
    print(json.dumps({**{k: v for k, v in audit.items() if k != 'source_polygon_issues'},
                      'source_polygon_issue_count': len(geometry_issues)}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
